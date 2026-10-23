"""
generate_packs.py
-----------------
Builds two synthetic test sets of quarterly reporting packs (Q2 2026, Apr-Jun):

  set_A_infrastructure/   5 fictional infrastructure portfolio companies
  set_B_asset_managers/   5 fictional asset-management firms

Every pack contains a metadata file (json or yaml) plus a mix of structured
(xlsx/csv/json), semi-structured (messy xlsx layouts) and unstructured
(pdf/docx/txt narrative) documents, with planted discrepancies.

Each set also gets:
  answer_key.csv      the correct extracted/derived values (ground truth)
  discrepancies.csv   every planted problem and what the pipeline should do

All companies are fictional. Numbers are invented.
"""

import csv, json, os, shutil, datetime as dt
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from docx import Document

ROOT = "/mnt/user-data/outputs/work/portfolio_test_packs_Q2_2026"
A = os.path.join(ROOT, "set_A_infrastructure")
B = os.path.join(ROOT, "set_B_asset_managers")

# FX to GBP. Flow items (P&L, cash flow, fund flows) use the quarter average;
# balances (net debt, AUM) use the closing rate on the balance date.
FX = {
    "GBP": {"avg_Q2_2026": 1.0, "close_2026-03-31": 1.0, "close_2026-06-30": 1.0},
    "EUR": {"avg_Q2_2026": 0.852, "close_2026-03-31": 0.861, "close_2026-06-30": 0.857},
    "USD": {"avg_Q2_2026": 0.742, "close_2026-03-31": 0.751, "close_2026-06-30": 0.738},
}

# --------------------------------------------------------------------------- helpers
def mk(path): os.makedirs(path, exist_ok=True); return path

def pdf(path, blocks):
    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=50, rightMargin=50)
    flow = []
    for kind, val in blocks:
        if kind == "h1": flow.append(Paragraph(val, st["Title"]))
        elif kind == "h2": flow.append(Paragraph(val, st["Heading2"]))
        elif kind == "p": flow.append(Paragraph(val, st["BodyText"]))
        elif kind == "small": flow.append(Paragraph(f"<font size=8>{val}</font>", st["BodyText"]))
        elif kind == "table":
            t = Table(val)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ]))
            flow.append(t)
        flow.append(Spacer(1, 8))
    doc.build(flow)

def docx(path, blocks):
    d = Document()
    for kind, val in blocks:
        if kind == "h1": d.add_heading(val, 0)
        elif kind == "h2": d.add_heading(val, 2)
        elif kind == "p": d.add_paragraph(val)
        elif kind == "bullets":
            for b in val: d.add_paragraph(b, style="List Bullet")
        elif kind == "table":
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for i, row in enumerate(val):
                for j, cell in enumerate(row): t.cell(i, j).text = str(cell)
    d.save(path)

def write_json(path, obj): open(path, "w").write(json.dumps(obj, indent=2, ensure_ascii=False))
def write_yaml(path, obj): open(path, "w").write(yaml.safe_dump(obj, sort_keys=False, allow_unicode=True))

def excel_serial(d): return (d - dt.date(1899, 12, 30)).days

# =========================================================================== SET A
# Ground truth per company: native currency, values in MILLIONS (calendar Q2 2026).
INFRA = {
 "A1": dict(name="Halden Flex Power Ltd", sector="Flexible power & storage (UK)", ccy="GBP",
     rev=42.1, rev_bud=44.0, rev_q1=38.5, ebitda=15.3, ebitda_bud=17.0, ebitda_q1=13.9,
     capex=22.4, capex_plan=25.0, ocf=12.9, net_debt=312.0, ltm_ebitda=56.0, cov_max=5.50,
     headcount=214, ltir=0.42),
 "A2": dict(name="Corvane Fibre Networks B.V.", sector="Long-haul fibre (NL/EU)", ccy="EUR",
     rev=61.8, rev_bud=60.5, rev_q1=59.2, ebitda=34.1, ebitda_bud=33.0, ebitda_q1=32.6,
     capex=28.7, capex_plan=None, ocf=30.2, net_debt=820.0, ltm_ebitda=131.0, cov_max=7.00,
     headcount=388, ltir=0.10),
 "A3": dict(name="Frostline Cold Logistics S.A.S.", sector="Cold storage (FR/ES)", ccy="EUR",
     rev=27.4, rev_bud=28.0, rev_q1=26.1, ebitda=7.9, ebitda_bud=8.6, ebitda_q1=7.2,
     capex=4.8, capex_plan=5.0, ocf=6.1, net_debt=96.0, ltm_ebitda=29.5, cov_max=4.00,
     headcount=640, ltir=1.9),
 "A4": dict(name="Meridian Bus Group Ltd", sector="Bus transport (UK)", ccy="GBP",
     rev=88.6, rev_bud=90.0, rev_q1=85.2, ebitda=11.2, ebitda_bud=12.4, ebitda_q1=10.8,
     capex=9.5, capex_plan=7.0, ocf=10.1, net_debt=64.0, ltm_ebitda=44.0, cov_max=3.00,
     headcount=None, ltir=7.8),
 "A5": dict(name="Verdan Recycling Holdings Ltd", sector="Waste & recycling (UK/IE)", ccy="GBP",
     rev=104.0, rev_bud=101.0, rev_q1=98.0, ebitda=18.2, ebitda_bud=19.0, ebitda_q1=17.0,
     capex=6.3, capex_plan=7.5, ocf=14.9, net_debt=210.0, ltm_ebitda=70.0, cov_max=3.50,
     headcount=2040, ltir=1.1),
}

def build_A1(d, c):
    # Messy Excel: merged headers, units in a title row, hidden D&A row, several tabs.
    wb = Workbook(); ws = wb.active; ws.title = "P&L"
    ws["A1"] = "HALDEN FLEX POWER LTD - MANAGEMENT ACCOUNTS"; ws["A1"].font = Font(bold=True)
    ws["A2"] = "£m unless stated | Period ended 30/06/2026"
    ws.merge_cells("B4:D4"); ws["B4"] = "Q2 2026 (Apr-Jun)"; ws["B4"].alignment = Alignment(horizontal="center")
    ws["E4"] = "Q1 2026"
    for col, h in zip("BCDE", ["Actual", "Budget", "Var", "Actual"]): ws[f"{col}5"] = h
    cos = round(c["rev"] * 0.48, 1); cos_b = round(c["rev_bud"] * 0.48, 1); cos_q1 = round(c["rev_q1"] * 0.48, 1)
    opex = round(c["rev"] - cos - c["ebitda"], 1); opex_b = round(c["rev_bud"] - cos_b - c["ebitda_bud"], 1)
    opex_q1 = round(c["rev_q1"] - cos_q1 - c["ebitda_q1"], 1)
    rows = [
        ("Revenue", c["rev"], c["rev_bud"], c["rev_q1"]),
        ("Cost of sales", -cos, -cos_b, -cos_q1),
        ("Gross margin", round(c["rev"] - cos, 1), round(c["rev_bud"] - cos_b, 1), round(c["rev_q1"] - cos_q1, 1)),
        ("Overheads", -opex, -opex_b, -opex_q1),
        ("EBITDA", c["ebitda"], c["ebitda_bud"], c["ebitda_q1"]),
        ("D&A", -6.8, -6.8, -6.6),
        ("Operating profit", round(c["ebitda"] - 6.8, 1), round(c["ebitda_bud"] - 6.8, 1), round(c["ebitda_q1"] - 6.6, 1)),
    ]
    for i, (lbl, a, b, q1) in enumerate(rows, start=6):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"], ws[f"E{i}"] = lbl, a, b, round(a - b, 1), q1
    ws.row_dimensions[11].hidden = True          # D&A row hidden
    ws["A15"] = "Note: budget re-phased in May; see board pack."
    cf = wb.create_sheet("Cash & Capex")
    cf.append(["£m", "Q2 2026 Actual", "Q2 2026 Plan"])
    cf.append(["Operating cash flow", c["ocf"], None])
    cf.append(["Capital expenditure", c["capex"], c["capex_plan"]])
    bs = wb.create_sheet("Debt")
    bs.append(["As at 30/06/2026", "£m"])
    bs.append(["Net debt", c["net_debt"]])
    bs.append(["EBITDA (LTM)", c["ltm_ebitda"]])
    bs.append(["Max leverage per facility agreement", "5.50x"])
    k = wb.create_sheet("KPIs")
    k.append(["KPI", "Q2 2026"]); k.append(["FTE", c["headcount"]]); k.append(["LTIR (per 200k hrs)", c["ltir"]])
    wb.save(os.path.join(d, "Halden_MA_Jun26.xlsx"))

    pdf(os.path.join(d, "Halden_Board_Pack_Q2_2026.pdf"), [
        ("h1", "Halden Flex Power - Board Pack"), ("small", "Board meeting 22/07/2026 | Quarter ended 30/06/2026"),
        ("h2", "CEO report"),
        ("p", "Q2 revenue of £42.1m was 4% below budget, driven by lower balancing-market prices in May. "
              "EBITDA of £15.3m was £1.7m behind budget. 30MW of battery capacity at Thornbury was energised in June."),
        ("h2", "Finance"),
        ("p", "Net debt closed the quarter at £312.0m. Leverage remains comfortably within the 5.5x covenant, "
              "and we do not anticipate any pressure on headroom over the next two quarters."),
        ("h2", "Risk register (extract)"),
        ("table", [["Risk", "Trend", "Owner"],
                   ["Grid connection delay - Fenwick site", "New", "COO"],
                   ["Balancing-market price volatility", "Worse", "CCO"],
                   ["Refinancing of 2027 facility", "Unchanged", "CFO"]]),
    ])
    write_json(os.path.join(d, "metadata.json"), {
        "company": c["name"], "reporting_period": "Q2 2026", "period_start": "01/04/2026", "period_end": "30/06/2026",
        "currency": "GBP", "units": "millions", "submitted": "21/07/2026", "version": "final",
        "prepared_by": "Group Finance", "files": ["Halden_MA_Jun26.xlsx", "Halden_Board_Pack_Q2_2026.pdf"]})

def build_A2(d, c):
    # Semicolon CSV export, thousands of EUR, decimal comma, DD.MM.YYYY, synonym label, missing capex plan.
    def k(x): return "" if x is None else f"{x*1000:.1f}".replace(".", ",")
    rows = [
        ["line_item", "period", "scenario", "value_kEUR", "as_of"],
        ["Revenue", "Q2-2026", "ACT", k(c["rev"]), "30.06.2026"],
        ["Revenue", "Q2-2026", "BUD", k(c["rev_bud"]), "30.06.2026"],
        ["Revenue", "Q1-2026", "ACT", k(c["rev_q1"]), "31.03.2026"],
        ["Adjusted operating result before D&A", "Q2-2026", "ACT", k(c["ebitda"]), "30.06.2026"],
        ["Adjusted operating result before D&A", "Q2-2026", "BUD", k(c["ebitda_bud"]), "30.06.2026"],
        ["Adjusted operating result before D&A", "Q1-2026", "ACT", k(c["ebitda_q1"]), "31.03.2026"],
        ["Cash from operations", "Q2-2026", "ACT", k(c["ocf"]), "30.06.2026"],
        ["Capex", "Q2-2026", "ACT", k(c["capex"]), "30.06.2026"],
        ["Capex", "Q2-2026", "BUD", k(c["capex_plan"]), "30.06.2026"],      # blank
        ["Net financial debt", "Q2-2026", "ACT", k(c["net_debt"]), "30.06.2026"],
        ["Adjusted operating result before D&A LTM", "Q2-2026", "ACT", k(c["ltm_ebitda"]), "30.06.2026"],
        ["Leverage covenant (max)", "Q2-2026", "", "7,00x", "30.06.2026"],
        ["Headcount (FTE)", "Q2-2026", "ACT", str(c["headcount"]), "30.06.2026"],
        ["LTIR", "Q2-2026", "ACT", "0,10", "30.06.2026"],
    ]
    with open(os.path.join(d, "corvane_finance_export_2026Q2.csv"), "w", newline="") as f:
        csv.writer(f, delimiter=";").writerows(rows)
    docx(os.path.join(d, "Corvane_Q2_commentary.docx"), [
        ("h1", "Corvane Fibre Networks - Q2 2026 management commentary"),
        ("p", "Revenue reached €61.8m in the second quarter, 2% ahead of budget, supported by two new wavelength contracts "
              "on the Amsterdam-Frankfurt route. Adjusted operating result before D&A was €34.1m."),
        ("p", "Capital expenditure was €28.7m, mainly the Hamburg ring extension. The re-phased capex plan is still with the board "
              "and is therefore not included in this pack."),
        ("h2", "Key risks"),
        ("bullets", ["Construction cost inflation on the Hamburg ring (worse)",
                     "Customer concentration: top 3 customers = 41% of revenue (unchanged)"]),
    ])
    write_yaml(os.path.join(d, "metadata.yaml"), {
        "entity_name": c["name"], "fiscal_period": "2026-Q2", "reporting_ccy": "EUR", "unit": "kEUR",
        "number_format": "decimal comma", "date_format": "DD.MM.YYYY", "submission_date": "24.07.2026",
        "status": "provisional", "contact": "controlling@corvane.example"})

def build_A3(d, c):
    # French PDF, capex unit error (48,0 instead of 4,8), Q1 restated; prior-quarter pack included.
    f = lambda x: f"{x:.1f}".replace(".", ",")
    pdf(os.path.join(d, "Frostline_Rapport_T2_2026.pdf"), [
        ("h1", "Frostline Cold Logistics - Rapport trimestriel T2 2026"),
        ("small", "Arrêté au 30 juin 2026 | Montants en millions d'euros (M€)"),
        ("table", [["Indicateur", "T2 2026 Réel", "T2 2026 Budget", "T1 2026 Réel (retraité)"],
                   ["Chiffre d'affaires", f(c["rev"]), f(c["rev_bud"]), f(c["rev_q1"])],
                   ["EBE (EBITDA)", f(c["ebitda"]), f(c["ebitda_bud"]), f(c["ebitda_q1"])],
                   ["Investissements (capex)", "48,0", f(c["capex_plan"]), "-"],
                   ["Flux de trésorerie opérationnel", f(c["ocf"]), "-", "-"],
                   ["Dette nette", f(c["net_debt"]), "-", "-"],
                   ["EBE 12 mois glissants", f(c["ltm_ebitda"]), "-", "-"],
                   ["Levier maximum (covenant)", "4,0x", "-", "-"],
                   ["Effectif (ETP)", str(c["headcount"]), "-", "-"],
                   ["Taux de fréquence des accidents (LTIR)", f(c["ltir"]), "-", "-"]]),
        ("small", "Le T1 2026 a été retraité suite au reclassement de certains frais de transport (impact CA : -0,8 M€ ; EBE : -0,4 M€)."),
        ("h2", "Commentaire de la direction"),
        ("p", "Le chiffre d'affaires du trimestre s'établit à 27,4 M€, légèrement inférieur au budget en raison d'un taux "
              "d'occupation plus faible sur le site de Valence. Les investissements sont restés conformes au plan."),
        ("h2", "Risques principaux"),
        ("table", [["Risque", "Tendance"], ["Coût de l'énergie (froid)", "Aggravé"],
                   ["Intégration de l'acquisition Ferrofrío", "Nouveau"]]),
    ])
    # Prior-quarter pack as originally reported (before restatement)
    pdf(os.path.join(d, "Frostline_Rapport_T1_2026_ORIGINAL.pdf"), [
        ("h1", "Frostline Cold Logistics - Rapport trimestriel T1 2026"),
        ("small", "Arrêté au 31 mars 2026 | M€"),
        ("table", [["Indicateur", "T1 2026 Réel"], ["Chiffre d'affaires", "26,9"], ["EBE (EBITDA)", "7,6"]]),
    ])
    write_json(os.path.join(d, "metadata.json"), {
        "societe": c["name"], "periode": "T2 2026", "date_arrete": "30 juin 2026", "devise": "EUR",
        "unite": "M€", "langue": "fr", "version": "définitive", "date_envoi": "28 juillet 2026"})

def build_A4(d, c):
    # Fiscal year Apr-Mar: calendar Q2 2026 = "Q1 FY27". Email with figures in prose; xlsx in £000 with Excel-serial dates.
    open(os.path.join(d, "email_from_CFO_Q1FY27.txt"), "w").write(
        "From: cfo@meridianbus.example\nTo: portfolio.monitoring@fund.example\nDate: Fri, 24 Jul 2026 17:42\n"
        "Subject: Q1 FY27 numbers\n\nHi both,\n\nQuick summary ahead of the full pack. Q1 FY27 (April-June) turnover came in at "
        "£88.6m against a plan of £90.0m; patronage on the Leeds network is still recovering after the timetable change. "
        "EBITDA was £11.2m (plan £12.4m). Capex was £9.5m versus £7.0m planned because we brought forward 14 electric buses.\n\n"
        "Net debt finished at £64.0m. Headcount data is delayed due to the HR system migration; I'll send it next week.\n\n"
        "One flag: LTIR jumped to 7.8 this quarter. We're checking whether that's real or a reporting error from the new system.\n\n"
        "Best,\nJ.\n")
    wb = Workbook(); ws = wb.active; ws.title = "Mgmt accounts"
    ws["A1"] = "Meridian Bus Group - £000"
    ws.append([]); ws.append(["Period end (date)", excel_serial(dt.date(2026, 6, 30)), excel_serial(dt.date(2026, 6, 30)),
                              excel_serial(dt.date(2026, 3, 31))])
    ws.append(["", "Q1 FY27 Actual", "Q1 FY27 Plan", "Q4 FY26 Actual"])
    t = lambda x: None if x is None else round(x * 1000)
    for lbl, a, b, p in [("Turnover", c["rev"], c["rev_bud"], c["rev_q1"]), ("EBITDA", c["ebitda"], c["ebitda_bud"], c["ebitda_q1"]),
                         ("Capex", c["capex"], c["capex_plan"], None), ("Operating cash flow", c["ocf"], None, None),
                         ("Net debt", c["net_debt"], None, None), ("EBITDA LTM", c["ltm_ebitda"], None, None)]:
        ws.append([lbl, t(a), t(b), t(p)])
    ws.append(["Covenant: max net debt / EBITDA", "3.0x"])
    ws.append(["Headcount", None])
    s = wb.create_sheet("Safety")
    s.append(["Quarter", "LTIR"])
    for q, v in [("Q2 FY26", 0.8), ("Q3 FY26", 0.7), ("Q4 FY26", 0.9), ("Q1 FY27", 7.8)]: s.append([q, v])
    wb.save(os.path.join(d, "Meridian_Q1FY27_MA.xlsx"))
    write_json(os.path.join(d, "metadata.json"), {
        "company": c["name"], "period_label": "Q1 FY27", "fiscal_year_start_month": 4,
        "period_covered": "2026-04-01 to 2026-06-30", "currency": "GBP", "units": "thousands (xlsx) / millions (email)",
        "received": "2026-07-24T17:42:00+01:00", "version": "draft - headcount outstanding"})

def build_A5(d, c):
    # Portal JSON export gives H1 YTD and Q1 only -> Q2 must be derived. US date format.
    q1_bud_rev, q1_bud_e, q1_capex, q1_capex_plan, q1_ocf = 97.0, 17.5, 5.0, 6.0, 13.0
    write_json(os.path.join(d, "verdan_portal_export.json"), {
        "entity": c["name"], "report_date": "06/30/2026", "currency": "GBP", "unit": "millions",
        "data": [
            {"metric": "Revenue", "period": "H1 2026 YTD", "actual": round(c["rev_q1"] + c["rev"], 1), "budget": q1_bud_rev + c["rev_bud"]},
            {"metric": "Revenue", "period": "Q1 2026", "actual": c["rev_q1"], "budget": q1_bud_rev},
            {"metric": "EBITDA", "period": "H1 2026 YTD", "actual": round(c["ebitda_q1"] + c["ebitda"], 1), "budget": q1_bud_e + c["ebitda_bud"]},
            {"metric": "EBITDA", "period": "Q1 2026", "actual": c["ebitda_q1"], "budget": q1_bud_e},
            {"metric": "Capex", "period": "H1 2026 YTD", "actual": round(q1_capex + c["capex"], 1), "budget": q1_capex_plan + c["capex_plan"]},
            {"metric": "Capex", "period": "Q1 2026", "actual": q1_capex, "budget": q1_capex_plan},
            {"metric": "Operating cash flow", "period": "H1 2026 YTD", "actual": round(q1_ocf + c["ocf"], 1), "budget": None},
            {"metric": "Operating cash flow", "period": "Q1 2026", "actual": q1_ocf, "budget": None},
            {"metric": "Net debt", "period": "as at 06/30/2026", "actual": c["net_debt"]},
            {"metric": "EBITDA LTM", "period": "as at 06/30/2026", "actual": c["ltm_ebitda"]},
            {"metric": "Max leverage covenant", "period": "facility", "actual": "3.5x"},
            {"metric": "Employees", "period": "as at 06/30/2026", "actual": c["headcount"]},
            {"metric": "LTIR", "period": "H1 2026 YTD", "actual": c["ltir"]},
        ]})
    docx(os.path.join(d, "Verdan_H1_2026_update.docx"), [
        ("h1", "Verdan Recycling - H1 2026 trading update"),
        ("p", "Second-quarter revenue grew 6% on the first quarter, driven by higher gate fees at the Avonmouth facility and "
              "stronger recovered-plastics prices. Year-to-date EBITDA is £35.2m."),
        ("p", "Leverage at 30 June was 3.0x against a 3.5x covenant. Headroom is tighter than we would like and we are "
              "prioritising cash generation in H2."),
        ("h2", "Risks"), ("bullets", ["Recovered-plastics price volatility (worse)", "Landfill tax change from April 2027 (new)"]),
    ])
    write_json(os.path.join(d, "metadata.json"), {
        "company": c["name"], "period": "H1 2026 (YTD)", "report_date": "06/30/2026", "date_format": "MM/DD/YYYY",
        "currency": "GBP", "units": "millions", "source_system": "investor portal export", "exported_at": "07/27/2026 09:15"})

# =========================================================================== SET B
AM = {
 "B1": dict(name="Ashcombe Capital Partners LLP", ccy="GBP", open_aum=48200, inflows=3100, outflows=2450, market=1180,
     close_aum=50030, revenue=61.3, avg_aum=49100, costs=42.8, out1=58, out3=64, out5=71, own_funds_cov=112, own_funds_min=120),
 "B2": dict(name="Lindqvist & Moor Asset Management S.A.", ccy="EUR", open_aum=21400, inflows=980, outflows=1640, market=410,
     close_aum=21150, revenue=23.9, avg_aum=21250, costs=17.6, out1=44, out3=51, out5=None),
 "B3": dict(name="Harbourline Investors Inc.", ccy="USD", open_aum=132500, inflows=7900, outflows=6300, market=3050,
     close_aum=137150, revenue=118.4, avg_aum=134800, costs=79.1, out1=61, out3=55, out5=67),
 "B4": dict(name="Kestrel Point Investment Management Ltd", ccy="GBP", open_aum=9850, inflows=410, outflows=690, market=215,
     close_aum=9785, revenue=14.6, avg_aum=9820, costs=12.1, out1=38, out3=47, out5=52),
 "B5": dict(name="Solenne Gestion S.A.", ccy="EUR", open_aum=15600, inflows=720, outflows=540, market=-260,
     close_aum=15520, revenue=15.1, avg_aum=15560, costs=11.7, out1=49, out3=62, out5=58),
}

def build_B1(d, c):
    wb = Workbook(); ws = wb.active; ws.title = "AUM rollforward"
    ws.append(["Ashcombe Capital Partners - AUM (£m)", None]); ws.append([])
    for r in [("Opening AUM 31/03/2026", c["open_aum"]), ("Gross sales", c["inflows"]), ("Redemptions", -c["outflows"]),
              ("Net new money", c["inflows"] - c["outflows"]), ("Market & FX", c["market"]), ("Closing AUM 30/06/2026", c["close_aum"]),
              ("Average AUM (quarter)", c["avg_aum"])]: ws.append(list(r))
    p = wb.create_sheet("Revenue & costs"); p.append(["£m", "Q2 2026"])
    p.append(["Management fees", round(c["revenue"] - 2.4, 1)]); p.append(["Performance fees", 2.4])
    p.append(["Total net revenue", c["revenue"]]); p.append(["Operating expenses", c["costs"]])
    f = wb.create_sheet("Performance"); f.append(["% of AUM above benchmark", "1y", "3y", "5y"])
    f.append(["Firm total", c["out1"] / 100, c["out3"] / 100, c["out5"] / 100])
    k = wb.create_sheet("Capital"); k.append(["Metric", "Value"])
    k.append(["Own funds / own funds requirement", "112%"]); k.append(["Board internal minimum", "120%"])
    wb.save(os.path.join(d, "Ashcombe_Q2_2026_KPIs.xlsx"))
    pdf(os.path.join(d, "Ashcombe_Board_Note_Q2.pdf"), [
        ("h1", "Ashcombe Capital Partners - Q2 board note"), ("small", "30 June 2026"),
        ("p", "A record quarter for gross sales (£3.1bn), led by the global credit strategy. Net new money was positive at £650m."),
        ("p", "The capital position remains robust and comfortably above all internal and regulatory minima."),
    ])
    write_json(os.path.join(d, "metadata.json"), {"firm": c["name"], "period": "Q2 2026", "period_end": "30/06/2026",
        "currency": "GBP", "aum_units": "millions", "revenue_units": "millions", "version": "final", "submitted": "20/07/2026"})

def build_B2(d, c):
    # Mixed units in one file: AUM in EUR bn, money flows in EUR m; 5y performance missing.
    rows = [["metric", "value", "unit", "period_end"],
            ["opening_aum", c["open_aum"] / 1000, "EUR bn", "2026-03-31"],
            ["gross_inflows", c["inflows"], "EUR m", "2026-06-30"],
            ["gross_outflows", c["outflows"], "EUR m", "2026-06-30"],
            ["market_effect", c["market"], "EUR m", "2026-06-30"],
            ["closing_aum", c["close_aum"] / 1000, "EUR bn", "2026-06-30"],
            ["average_aum", c["avg_aum"] / 1000, "EUR bn", "2026-06-30"],
            ["net_revenue", c["revenue"], "EUR m", "2026-06-30"],
            ["operating_costs", c["costs"], "EUR m", "2026-06-30"],
            ["pct_aum_outperforming_1y", c["out1"], "%", "2026-06-30"],
            ["pct_aum_outperforming_3y", c["out3"], "%", "2026-06-30"],
            ["pct_aum_outperforming_5y", "", "%", "2026-06-30"]]
    with open(os.path.join(d, "lindqvist_moor_q2_2026.csv"), "w", newline="") as f: csv.writer(f).writerows(rows)
    open(os.path.join(d, "commentary.txt"), "w").write(
        "Q2 2026 commentary - Lindqvist & Moor\n\nNet outflows of EUR 660m were concentrated in European equity income, where two "
        "institutional mandates were lost. Fixed income saw modest net inflows. Five-year track records are being re-calculated "
        "following the fund merger in March and are therefore omitted this quarter.\n")
    write_yaml(os.path.join(d, "metadata.yaml"), {"entity": c["name"], "quarter": "2026Q2", "period_end": "2026-06-30",
        "ccy": "EUR", "units": "mixed - see unit column", "status": "final"})

def build_B3(d, c):
    # US fiscal year ending Sept: calendar Q2 2026 = fiscal Q3 2026. Rollforward does not tie (closing typo +500).
    n = lambda x: f"{x:,.0f}"
    pdf(os.path.join(d, "Harbourline_FQ3_2026_Investor_Update.pdf"), [
        ("h1", "Harbourline Investors - Fiscal Q3 2026 Update"), ("small", "Quarter ended 06/30/2026 | USD millions"),
        ("table", [["AUM rollforward", "USD m"], ["Beginning AUM (03/31/2026)", n(c["open_aum"])], ["Inflows", n(c["inflows"])],
                   ["Outflows", f"({n(c['outflows'])})"], ["Net flows", n(c["inflows"] - c["outflows"])],
                   ["Market appreciation", n(c["market"])], ["Ending AUM (06/30/2026)", "137,650"], ["Average AUM", n(c["avg_aum"])]]),
        ("table", [["Financials", "FQ3 2026"], ["Revenues", f"{c['revenue']:.1f}"], ["Operating expenses", f"{c['costs']:.1f}"]]),
        ("table", [["% of AUM outperforming benchmark", "1-yr", "3-yr", "5-yr"], ["All strategies", f"{c['out1']}%", f"{c['out3']}%", f"{c['out5']}%"]]),
        ("p", "Positive net flows for the fifth consecutive quarter, led by our US large-cap growth franchise."),
    ])
    write_json(os.path.join(d, "metadata.json"), {"firm": c["name"], "period_label": "FQ3 2026", "fiscal_year_end": "09-30",
        "period_end": "06/30/2026", "date_format": "MM/DD/YYYY", "currency": "USD", "units": "millions"})

def build_B4(d, c):
    # Narrative-heavy docx; outflows 10x outlier in the table; text contradicts the numbers.
    docx(os.path.join(d, "Kestrel_Point_Q2_letter.docx"), [
        ("h1", "Kestrel Point - Quarterly letter to the board, Q2 2026"),
        ("p", f"Assets under management ended the quarter at £{c['close_aum']:,}m, from £{c['open_aum']:,}m at the end of March. "
              f"Gross sales of £{c['inflows']}m were offset by redemptions of £{c['outflows']}m, mostly from one pension client. "
              f"Markets added £{c['market']}m."),
        ("p", "We continue to see positive net flow momentum across the franchise, and the pipeline for H2 is encouraging."),
        ("p", f"Net revenue for the quarter was £{c['revenue']}m and operating costs were £{c['costs']}m, reflecting continued "
              f"investment in the new client-reporting platform. Average AUM was £{c['avg_aum']:,}m."),
        ("p", f"The share of AUM ahead of benchmark is {c['out1']}% over one year, {c['out3']}% over three years and {c['out5']}% over five."),
        ("h2", "Summary table (£m)"),
        ("table", [["Opening AUM", "Sales", "Redemptions", "Market", "Closing AUM"],
                   [f"{c['open_aum']:,}", c["inflows"], "6,900", c["market"], f"{c['close_aum']:,}"]]),
    ])
    write_json(os.path.join(d, "metadata.json"), {"firm": c["name"], "period": "Q2 2026", "as_at": "30 June 2026",
        "currency": "GBP", "units": "millions", "document_type": "board letter (narrative)", "version": "final"})

def build_B5(d, c):
    # French PDF; opening AUM restated (15,900 -> 15,600); prior pack included.
    n = lambda x: f"{x:,.0f}".replace(",", " ")
    pdf(os.path.join(d, "Solenne_Reporting_T2_2026.pdf"), [
        ("h1", "Solenne Gestion - Reporting trimestriel T2 2026"), ("small", "Au 30/06/2026 | en millions d'euros"),
        ("table", [["Indicateur", "T2 2026"],
                   ["Encours sous gestion au 31/03/2026 (retraité)", n(c["open_aum"])],
                   ["Collecte brute", n(c["inflows"])], ["Décollecte", n(c["outflows"])],
                   ["Collecte nette", n(c["inflows"] - c["outflows"])], ["Effet marché", n(c["market"])],
                   ["Encours sous gestion au 30/06/2026", n(c["close_aum"])], ["Encours moyen", n(c["avg_aum"])],
                   ["Commissions nettes", f"{c['revenue']:.1f}".replace(".", ",")],
                   ["Charges d'exploitation", f"{c['costs']:.1f}".replace(".", ",")]]),
        ("table", [["% des encours surperformant l'indice", "1 an", "3 ans", "5 ans"], ["Total", f"{c['out1']} %", f"{c['out3']} %", f"{c['out5']} %"]]),
        ("small", "Encours au 31/03/2026 retraité de 15 900 M€ à 15 600 M€ suite à l'exclusion d'un mandat de conseil (non géré)."),
    ])
    pdf(os.path.join(d, "Solenne_Reporting_T1_2026_ORIGINAL.pdf"), [
        ("h1", "Solenne Gestion - Reporting trimestriel T1 2026"), ("small", "Au 31/03/2026 | M€"),
        ("table", [["Indicateur", "T1 2026"], ["Encours sous gestion au 31/03/2026", "15 900"]]),
    ])
    write_json(os.path.join(d, "metadata.json"), {"societe": c["name"], "periode": "T2 2026", "date_arrete": "30/06/2026",
        "devise": "EUR", "unite": "M€", "langue": "fr"})

# =========================================================================== answer keys
def infra_answer_rows():
    out = []
    for cid, c in INFRA.items():
        f_avg, f_close = FX[c["ccy"]]["avg_Q2_2026"], FX[c["ccy"]]["close_2026-06-30"]
        def add(metric, scen, val, kind, note=""):
            gbp = "" if val is None or kind == "ratio" else round(val * (f_close if kind == "balance" else f_avg), 2)
            out.append(["A", cid, c["name"], "2026-Q2", metric, scen,
                        "" if val is None else round(val, 4), c["ccy"] if kind != "ratio" else "",
                        "m" if kind in ("flow", "balance") else ("" if kind == "ratio" else kind), gbp, note])
        add("revenue", "actual", c["rev"], "flow"); add("revenue", "budget", c["rev_bud"], "flow"); add("revenue", "prior_quarter", c["rev_q1"], "flow")
        add("ebitda", "actual", c["ebitda"], "flow"); add("ebitda", "budget", c["ebitda_bud"], "flow"); add("ebitda", "prior_quarter", c["ebitda_q1"], "flow")
        add("capex", "actual", c["capex"], "flow"); add("capex", "plan", c["capex_plan"], "flow", "missing in pack" if c["capex_plan"] is None else "")
        add("operating_cash_flow", "actual", c["ocf"], "flow"); add("net_debt", "actual", c["net_debt"], "balance")
        add("ebitda_ltm", "actual", c["ltm_ebitda"], "flow"); add("covenant_max_leverage", "facility", c["cov_max"], "ratio")
        add("headcount", "actual", c["headcount"], "FTE", "missing in pack" if c["headcount"] is None else "")
        add("ltir", "actual", c["ltir"], "per 200k hrs",
            {"A4": "anomaly vs 0.7-0.9 history - flag, unconfirmed", "A5": "reported as H1 YTD, not Q2"}.get(cid, ""))
        lev = c["net_debt"] / c["ltm_ebitda"]
        add("ebitda_margin_pct", "actual", c["ebitda"] / c["rev"] * 100, "ratio")
        add("revenue_vs_budget_pct", "actual", (c["rev"] / c["rev_bud"] - 1) * 100, "ratio")
        add("ebitda_vs_budget_pct", "actual", (c["ebitda"] / c["ebitda_bud"] - 1) * 100, "ratio")
        add("ebitda_qoq_pct", "actual", (c["ebitda"] / c["ebitda_q1"] - 1) * 100, "ratio")
        add("cash_conversion_pct", "actual", c["ocf"] / c["ebitda"] * 100, "ratio")
        add("capex_vs_plan_pct", "actual", None if c["capex_plan"] is None else (c["capex"] / c["capex_plan"] - 1) * 100, "ratio",
            "not computable - plan missing" if c["capex_plan"] is None else "")
        add("leverage_x", "actual", lev, "ratio")
        add("covenant_headroom_x", "actual", c["cov_max"] - lev, "ratio",
            "BREACH" if lev > c["cov_max"] else ("tight (<0.5x)" if c["cov_max"] - lev <= 0.5 else ""))
    return out

def am_answer_rows():
    out = []
    for cid, c in AM.items():
        fx = FX[c["ccy"]]
        def add(metric, val, kind, note="", rate_key=None):
            gbp = "" if val is None or kind == "ratio" else round(val * fx[rate_key], 2)
            out.append(["B", cid, c["name"], "2026-Q2", metric, "actual", "" if val is None else round(val, 4),
                        c["ccy"] if kind != "ratio" else "", "m" if kind != "ratio" else "", gbp, note])
        net = c["inflows"] - c["outflows"]
        add("opening_aum", c["open_aum"], "bal", "restated (prior pack said 15,900)" if cid == "B5" else "", "close_2026-03-31")
        add("gross_inflows", c["inflows"], "flow", rate_key="avg_Q2_2026")
        add("gross_outflows", c["outflows"], "flow", "pack table shows 6,900 (outlier)" if cid == "B4" else "", "avg_Q2_2026")
        add("net_flows", net, "flow", rate_key="avg_Q2_2026")
        add("market_movement", c["market"], "flow", rate_key="avg_Q2_2026")
        add("closing_aum", c["close_aum"], "bal", "pack shows 137,650 - rollforward does not tie" if cid == "B3" else "", "close_2026-06-30")
        add("average_aum", c["avg_aum"], "bal", rate_key="avg_Q2_2026")
        add("net_revenue", c["revenue"], "flow", rate_key="avg_Q2_2026")
        add("operating_costs", c["costs"], "flow", rate_key="avg_Q2_2026")
        add("aum_growth_pct", (c["close_aum"] / c["open_aum"] - 1) * 100, "ratio")
        add("organic_growth_pct_quarterly", net / c["open_aum"] * 100, "ratio")
        add("organic_growth_pct_annualised", net / c["open_aum"] * 400, "ratio")
        add("market_contribution_pct", c["market"] / c["open_aum"] * 100, "ratio")
        add("fee_rate_bps_annualised", c["revenue"] * 4 / c["avg_aum"] * 10000, "ratio")
        add("operating_margin_pct", (c["revenue"] - c["costs"]) / c["revenue"] * 100, "ratio")
        add("cost_income_ratio_pct", c["costs"] / c["revenue"] * 100, "ratio")
        add("pct_aum_outperforming_1y", c["out1"], "ratio"); add("pct_aum_outperforming_3y", c["out3"], "ratio")
        add("pct_aum_outperforming_5y", c["out5"], "ratio", "missing in pack" if c["out5"] is None else "")
        if "own_funds_cov" in c:
            add("own_funds_coverage_pct", c["own_funds_cov"], "ratio", f"BELOW internal minimum {c['own_funds_min']}%")
    return out

HEADER = ["set", "company_id", "company", "period_calendar", "metric", "scenario", "expected_value_native",
          "native_currency", "unit", "expected_value_gbp_m", "note"]

DISC_A = [
 ["A1", "Halden_MA_Jun26.xlsx", "P&L tab", "semi-structured layout", "Merged header cells, units only in a title row, hidden D&A row, free-text note in A15", "Read units from title; ignore hidden row only if it is not needed"],
 ["A1", "Debt tab + board pack", "covenant", "threshold breach", "Leverage 312.0/56.0 = 5.57x vs 5.50x max", "Flag BREACH"],
 ["A1", "Halden_Board_Pack_Q2_2026.pdf", "Finance section", "narrative contradicts numbers", "Says leverage 'comfortably within' covenant", "Flag contradiction; numbers win"],
 ["A1", "metadata.json", "dates", "date format", "DD/MM/YYYY", "Parse as day-first"],
 ["A2", "corvane_finance_export_2026Q2.csv", "whole file", "format", "Semicolon delimiter, decimal comma, values in kEUR", "Divide by 1000; parse decimal comma"],
 ["A2", "corvane_finance_export_2026Q2.csv", "EBITDA rows", "synonym label", "EBITDA reported as 'Adjusted operating result before D&A'", "Map to ebitda via KPI dictionary"],
 ["A2", "corvane_finance_export_2026Q2.csv", "Capex BUD row", "missing value", "Capex plan blank (explained in commentary)", "Leave null; flag; capex_vs_plan not computable"],
 ["A2", "all", "currency", "currency", "EUR", "Convert to GBP: flows at Q2 average, net debt at 30/06 close"],
 ["A2", "metadata.yaml", "dates / keys", "date format + different metadata keys", "DD.MM.YYYY; keys differ from other packs (reporting_ccy, unit)", "Normalise metadata keys"],
 ["A3", "Frostline_Rapport_T2_2026.pdf", "whole file", "language", "French labels (Chiffre d'affaires, EBE, Dette nette), decimal comma", "Map French labels"],
 ["A3", "Frostline_Rapport_T2_2026.pdf", "capex", "outlier / unit error", "Capex shown 48,0 vs plan 5,0 and commentary 'in line with plan' -> true 4.8", "Flag outlier; do not silently correct"],
 ["A3", "T2 vs T1 ORIGINAL pdf", "Q1 revenue/EBITDA", "restatement", "Q1 revenue 26.9 -> 26.1, EBITDA 7.6 -> 7.2", "Use restated figures for QoQ; flag restatement"],
 ["A3", "metadata.json", "dates", "date format", "'30 juin 2026' (French text date)", "Parse localised date"],
 ["A4", "metadata.json + xlsx", "period", "period definition", "Fiscal year starts April: 'Q1 FY27' = calendar Q2 2026; prior column 'Q4 FY26' = calendar Q1 2026", "Map fiscal to calendar quarter"],
 ["A4", "Meridian_Q1FY27_MA.xlsx", "whole sheet", "units + dates", "Values in £000; period-end dates stored as Excel serial numbers", "Divide by 1000; convert serial dates"],
 ["A4", "email_from_CFO_Q1FY27.txt", "body", "unstructured source", "Key numbers only in email prose, in £m", "Extract from text; reconcile with xlsx"],
 ["A4", "xlsx Headcount row", "headcount", "missing value", "Headcount blank (HR migration)", "Leave null; flag"],
 ["A4", "Safety tab", "LTIR", "anomaly", "LTIR 7.8 vs 0.7-0.9 in prior quarters; CFO unsure if real", "Flag anomaly for follow-up"],
 ["A4", "Cash/capex", "capex", "variance", "Capex 36% over plan (explained: electric buses brought forward)", "Report variance with explanation"],
 ["A5", "verdan_portal_export.json", "P&L/cash items", "period definition", "Only H1 YTD and Q1 given", "Derive Q2 = H1 YTD - Q1"],
 ["A5", "verdan_portal_export.json", "dates", "date format", "MM/DD/YYYY (US)", "Parse month-first"],
 ["A5", "covenant", "leverage", "threshold near-breach", "3.0x vs 3.5x max (0.5x headroom)", "Flag tight headroom"],
 ["A5", "LTIR", "period", "period definition", "LTIR given H1 YTD, not quarterly", "Label as H1 YTD; do not treat as Q2"],
]
DISC_B = [
 ["B1", "Ashcombe_Q2_2026_KPIs.xlsx", "Capital tab", "threshold breach", "Own funds coverage 112% vs board minimum 120%", "Flag BREACH"],
 ["B1", "Ashcombe_Board_Note_Q2.pdf", "body", "narrative contradicts numbers", "Says capital 'comfortably above all internal minima'", "Flag contradiction"],
 ["B1", "Performance tab", "values", "format", "Percentages stored as decimals (0.58)", "Convert to %"],
 ["B1", "AUM rollforward", "labels", "synonym label", "'Net new money', 'Gross sales', 'Redemptions'", "Map to net_flows / gross_inflows / gross_outflows"],
 ["B2", "lindqvist_moor_q2_2026.csv", "AUM rows", "mixed units", "AUM in EUR bn, flows/revenue in EUR m in the same file", "Normalise to millions"],
 ["B2", "lindqvist_moor_q2_2026.csv", "5y performance", "missing value", "Blank; commentary explains fund merger", "Leave null; flag"],
 ["B2", "all", "currency", "currency", "EUR", "Convert to GBP"],
 ["B3", "Harbourline_FQ3_2026_Investor_Update.pdf", "AUM rollforward", "does not tie", "132,500 + 1,600 + 3,050 = 137,150, but ending AUM printed as 137,650", "Flag; components imply 137,150"],
 ["B3", "metadata.json", "period", "period definition", "Fiscal year ends Sept: 'FQ3 2026' = calendar Q2 2026", "Map fiscal to calendar"],
 ["B3", "all", "currency + dates", "currency / date format", "USD; MM/DD/YYYY", "Convert to GBP; parse month-first"],
 ["B4", "Kestrel_Point_Q2_letter.docx", "summary table", "outlier", "Redemptions 6,900 in table vs 690 in text; only 690 ties the rollforward", "Flag; use value that ties"],
 ["B4", "Kestrel_Point_Q2_letter.docx", "paragraph 2", "narrative contradicts numbers", "Claims 'positive net flow momentum' while net flows are -280", "Flag contradiction"],
 ["B4", "whole pack", "structure", "unstructured source", "All numbers in prose apart from one table", "Extract from text"],
 ["B5", "Solenne_Reporting_T2_2026.pdf", "whole file", "language", "French labels (encours, collecte nette, effet marché, commissions)", "Map French labels"],
 ["B5", "T2 vs T1 ORIGINAL pdf", "opening AUM", "restatement", "Opening AUM restated 15,900 -> 15,600", "Use restated; flag"],
 ["B5", "Solenne_Reporting_T2_2026.pdf", "numbers", "format", "Space as thousands separator, decimal comma", "Parse locale"],
]
DISC_HEADER = ["company_id", "file", "location", "discrepancy_type", "what_was_planted", "expected_pipeline_behaviour"]

def write_csv(path, header, rows):
    with open(path, "w", newline="") as f: w = csv.writer(f); w.writerow(header); w.writerows(rows)

# =========================================================================== main
if __name__ == "__main__":
    if os.path.exists(ROOT): shutil.rmtree(ROOT)
    builders_A = {"A1": build_A1, "A2": build_A2, "A3": build_A3, "A4": build_A4, "A5": build_A5}
    builders_B = {"B1": build_B1, "B2": build_B2, "B3": build_B3, "B4": build_B4, "B5": build_B5}
    for cid, fn in builders_A.items():
        fn(mk(os.path.join(A, "packs", f"{cid}_{INFRA[cid]['name'].split()[0].lower()}")), INFRA[cid])
    for cid, fn in builders_B.items():
        fn(mk(os.path.join(B, "packs", f"{cid}_{AM[cid]['name'].split()[0].lower()}")), AM[cid])
    write_csv(os.path.join(A, "answer_key.csv"), HEADER, infra_answer_rows())
    write_csv(os.path.join(B, "answer_key.csv"), HEADER, am_answer_rows())
    write_csv(os.path.join(A, "discrepancies.csv"), DISC_HEADER, DISC_A)
    write_csv(os.path.join(B, "discrepancies.csv"), DISC_HEADER, DISC_B)
    write_csv(os.path.join(ROOT, "fx_rates_to_GBP.csv"), ["currency", "avg_Q2_2026", "close_2026-03-31", "close_2026-06-30"],
              [[k, v["avg_Q2_2026"], v["close_2026-03-31"], v["close_2026-06-30"]] for k, v in FX.items()])
    print("done")
