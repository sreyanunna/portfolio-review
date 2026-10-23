"""Ingest: read every file in a pack, as-is.

Every reader returns the same shape:
    {"file": name, "type": ..., "tables": [...], "text": [...], "data": ...}

- tables: list of {"name", "rows", "hidden_rows"}; rows are lists of raw cell values
- text:   list of paragraphs / page text
- data:   the parsed object (json / yaml only)

Nothing is cleaned here: "61800,0", "5.50x" and 46203 come through unchanged.
"""

import csv
import json
from pathlib import Path

import docx
import openpyxl
import pdfplumber
import yaml


def _result(path, kind, tables=None, text=None, data=None):
    return {"file": path.name, "type": kind, "tables": tables or [], "text": text or [], "data": data}


def read_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        hidden = [r for r, dim in ws.row_dimensions.items() if dim.hidden]
        tables.append({"name": ws.title, "rows": rows, "hidden_rows": hidden})
    return _result(path, "excel", tables=tables)


def read_csv(path):
    text = path.read_text(encoding="utf-8-sig")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
    rows = list(csv.reader(text.splitlines(), dialect))
    return _result(path, "csv", tables=[{"name": path.stem, "rows": rows, "hidden_rows": []}])


def read_json(path):
    return _result(path, "json", data=json.loads(path.read_text(encoding="utf-8")))


def read_yaml(path):
    return _result(path, "yaml", data=yaml.safe_load(path.read_text(encoding="utf-8")))


def read_pdf(path):
    tables, text = [], []
    with pdfplumber.open(path) as pdf:
        for p, page in enumerate(pdf.pages, start=1):
            for t, rows in enumerate(page.extract_tables(), start=1):
                tables.append({"name": f"page{p}_table{t}", "rows": rows, "hidden_rows": []})
            text.append(page.extract_text() or "")
    return _result(path, "pdf", tables=tables, text=text)


def read_docx(path):
    doc = docx.Document(str(path))
    text = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = [
        {"name": f"table{t}", "rows": [[c.text for c in row.cells] for row in table.rows],
         "hidden_rows": []}
        for t, table in enumerate(doc.tables, start=1)
    ]
    return _result(path, "docx", tables=tables, text=text)


def read_txt(path):
    text = path.read_text(encoding="utf-8-sig")
    return _result(path, "txt", text=[p.strip() for p in text.split("\n\n") if p.strip()])


READERS = {
    ".xlsx": read_excel,
    ".csv": read_csv,
    ".json": read_json,
    ".yaml": read_yaml,
    ".yml": read_yaml,
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".txt": read_txt,
}


def ingest_file(path):
    path = Path(path)
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        return _result(path, "unsupported")
    return reader(path)


def ingest_pack(pack_dir):
    """All files in one company's pack folder, e.g. data/set_A_infrastructure/packs/A1_halden."""
    files = sorted(p for p in Path(pack_dir).iterdir() if p.is_file() and not p.name.startswith("."))
    return [ingest_file(p) for p in files]


def ingest_set(set_dir):
    """{pack folder name: [file results]} for every pack under <set_dir>/packs."""
    packs = sorted(p for p in (Path(set_dir) / "packs").iterdir() if p.is_dir())
    return {pack.name: ingest_pack(pack) for pack in packs}
